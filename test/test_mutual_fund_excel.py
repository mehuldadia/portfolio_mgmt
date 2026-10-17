"""Tests for the mutual-fund workbook template and parser."""

from io import BytesIO
import os
import sys
import unittest

from openpyxl import Workbook, load_workbook

sys.path.insert(
    0,
    os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "src"),
)

from modules.mutual_fund_excel import (
    create_mutual_fund_template,
    read_mutual_fund_transactions,
)
from modules.mutual_fund_transactions import (
    dividend_records_to_dataframe,
    trade_records_to_dataframe,
)


EXPECTED_COLUMNS = [
    "MF_NAME",
    "INVESTOR_NAME",
    "PAN",
    "FOLIO_NUMBER",
    "PRODUCT_CODE",
    "SCHEME_NAME",
    "Scheme Type",
    "TRADE_DATE",
    "TRANSACTION_TYPE",
    "DIVIDEND_RATE",
    "AMOUNT",
    "UNITS",
    "PRICE",
    "BROKER",
]


class MutualFundExcelTests(unittest.TestCase):
    def test_template_has_requested_columns_in_order(self):
        workbook = load_workbook(create_mutual_fund_template())

        self.assertEqual(workbook.sheetnames, ["Transactions"])
        self.assertEqual(
            [cell.value for cell in workbook["Transactions"][1]], EXPECTED_COLUMNS
        )

    def test_reader_parses_trades_and_dividends_with_metadata(self):
        workbook = load_workbook(create_mutual_fund_template())
        sheet = workbook["Transactions"]
        sheet.append(
            [
                "ABC Mutual Fund",
                "Test Investor",
                "ABCDE1234F",
                "FOLIO-1",
                "P001",
                "Balanced Fund",
                "Equity",
                "2025-01-02",
                "PURCHASE",
                None,
                100,
                10,
                10,
                "Broker X",
            ]
        )
        sheet.append(
            [
                "ABC Mutual Fund",
                "Test Investor",
                "ABCDE1234F",
                "FOLIO-1",
                "P001",
                "Balanced Fund",
                "Equity",
                "2025-02-02",
                "DIVIDEND",
                1.5,
                5,
                None,
                None,
                "Broker X",
            ]
        )
        stream = BytesIO()
        workbook.save(stream)
        stream.seek(0)

        trades, dividends = read_mutual_fund_transactions(stream)

        self.assertEqual(len(trades), 1)
        self.assertEqual(trades[0]["Type"], "BUY")
        self.assertEqual(trades[0]["Amount"], 100)
        self.assertEqual(trades[0]["Investor_Name"], "Test Investor")
        self.assertEqual(trades[0]["PAN"], "ABCDE1234F")
        self.assertEqual(trades[0]["Folio_Number"], "FOLIO-1")
        self.assertEqual(trades[0]["Scheme_Name"], "Balanced Fund")
        self.assertEqual(trades[0]["Broker"], "Broker X")
        self.assertEqual(len(dividends), 1)
        self.assertEqual(dividends[0]["Amount"], 5)
        self.assertEqual(dividends[0]["Dividend_Rate"], 1.5)
        self.assertEqual(dividends[0]["Scheme_Type"], "Equity")
        self.assertEqual(trade_records_to_dataframe(trades).iloc[0]["PAN"], "ABCDE1234F")
        self.assertEqual(
            dividend_records_to_dataframe(dividends).iloc[0]["Investor_Name"],
            "Test Investor",
        )

    def test_reader_keeps_support_for_previous_two_sheet_format(self):
        workbook = Workbook()
        trades = workbook.active
        trades.title = "Transactions"
        trades.append(["Txn_Date", "Fund_Name", "ISIN", "Type", "Quantity", "Price", "Fees"])
        trades.append(["2025-01-02", "Legacy Fund", "INF100", "Purchase", 4, 5, 0])
        dividends = workbook.create_sheet("Dividends")
        dividends.append(["Date", "Fund_Name", "ISIN", "Amount", "Notes"])
        dividends.append(["2025-02-02", "Legacy Fund", "INF100", 3, "Old format"])
        stream = BytesIO()
        workbook.save(stream)
        stream.seek(0)

        trade_records, dividend_records = read_mutual_fund_transactions(stream)

        self.assertEqual(len(trade_records), 1)
        self.assertEqual(trade_records[0]["Type"], "BUY")
        self.assertEqual(len(dividend_records), 1)
        self.assertEqual(dividend_records[0]["Amount"], 3)


if __name__ == "__main__":
    unittest.main()
