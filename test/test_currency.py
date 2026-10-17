"""Tests for INR display formatting."""

import os
import sys
import tempfile
import unittest

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(
    0,
    os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "src"),
)

from modules.currency import format_inr
from modules.excel_exporter import ExcelExporter


class CurrencyFormattingTests(unittest.TestCase):
    def test_formats_indian_grouping_and_sign(self):
        self.assertEqual(format_inr(1234567.89), "₹12,34,567.89")
        self.assertEqual(format_inr(-1234.5), "-₹1,234.50")

    def test_empty_value_is_zero(self):
        self.assertEqual(format_inr(None), "₹0.00")

    def test_excel_ledger_keeps_values_numeric_with_inr_format(self):
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as output_dir:
            exporter = ExcelExporter(output_dir)
            path = exporter.export_mutual_fund_ledger(
                pd.DataFrame(
                    [{
                        "Txn_Date": "2025-01-01",
                        "Symbol": "P1",
                        "ISIN_Name": "Fund",
                        "Type": "BUY",
                        "Quantity": 1,
                        "Price": 1234.5,
                        "Amount": 1234.5,
                        "Fees": 0,
                        "Total_Value": 1234.5,
                    }]
                ),
                pd.DataFrame(
                    [{"Date": "2025-01-02", "Fund_Name": "Fund", "ISIN": "P1", "Amount": 12, "Notes": ""}]
                ),
            )
            workbook = load_workbook(path, read_only=True)
            try:
                sheet = workbook["MF_Trades"]
                columns = {cell.value: cell.column for cell in sheet[1]}
                price_cell = sheet.cell(2, columns["Price"])
                self.assertEqual(price_cell.value, 1234.5)
                self.assertIn("₹", price_cell.number_format)
                sheet = workbook["Dividends"]
                columns = {cell.value: cell.column for cell in sheet[1]}
                self.assertIn("₹", sheet.cell(2, columns["Amount"]).number_format)
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()
